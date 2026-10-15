from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.enum.shapes import PP_PLACEHOLDER
from pathlib import Path

OUT = Path(__file__).parent
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

BG = RGBColor(9, 16, 28); PANEL = RGBColor(18, 30, 46); PANEL2 = RGBColor(24, 39, 58)
WHITE = RGBColor(242, 247, 250); MUTED = RGBColor(161, 180, 194); CYAN = RGBColor(50, 205, 220)
BLUE = RGBColor(64, 139, 255); RED = RGBColor(235, 83, 76); ORANGE = RGBColor(245, 163, 62)
GREEN = RGBColor(91, 201, 132); GRID = RGBColor(47, 67, 86); PURPLE = RGBColor(164, 125, 241)
FONT = 'Aptos'; MONO = 'Cascadia Mono'

def rect(slide, x,y,w,h, fill=PANEL, line=None, radius=False):
    s=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb=fill
    s.line.color.rgb = line or fill
    if radius: s.adjustments[0]=0.08
    return s
def line(slide,x1,y1,x2,y2,color=GRID,width=1,dash=None):
    s=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)); s.line.color.rgb=color; s.line.width=Pt(width)
    if dash: s.line.dash_style = dash
    return s
def text(slide, x,y,w,h, value, size=16, color=WHITE, bold=False, font=FONT, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.06):
    tb=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf=tb.text_frame; tf.clear(); tf.word_wrap=True; tf.margin_left=Inches(margin); tf.margin_right=Inches(margin); tf.margin_top=Inches(margin); tf.margin_bottom=Inches(margin); tf.vertical_anchor=valign
    p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=value; r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color
    return tb
def add_text_runs(slide,x,y,w,h,runs,size=16,align=PP_ALIGN.LEFT):
    tb=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame; tf.clear(); tf.word_wrap=True; tf.margin_left=Inches(.06); tf.margin_right=Inches(.06)
    p=tf.paragraphs[0]; p.alignment=align
    for val,col,bold in runs:
        r=p.add_run(); r.text=val; r.font.name=FONT; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=col
    return tb
def base(title, section, num):
    s=prs.slides.add_slide(blank); rect(s,0,0,13.333,7.5,BG); rect(s,.42,.38,.06,.55,CYAN)
    text(s,.62,.35,9.8,.3,section.upper(),9,CYAN,True); text(s,.62,.72,12,.55,title,27,WHITE,True)
    text(s,.62,7.12,6,.18,'RESCUESWARM  /  AUTONOMOUS EMERGENCY RESPONSE',7,MUTED,True,font=MONO)
    text(s,12.3,7.08,.45,.2,f'{num:02d}',9,MUTED,True,font=MONO,align=PP_ALIGN.RIGHT)
    return s
def pill(s,x,y,w,label,color): rect(s,x,y,w,.28,color,color,True); text(s,x,y+.01,w,.2,label,8,BG,True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
def card(s,x,y,w,h,title_,body='',accent=CYAN, big=None):
    rect(s,x,y,w,h,PANEL,GRID,True); rect(s,x,y,.05,h,accent)
    text(s,x+.18,y+.14,w-.3,.22,title_.upper(),9,accent,True,font=MONO)
    if big: text(s,x+.18,y+.44,w-.3,.46,big,25,WHITE,True)
    if body: text(s,x+.18,y+(1.02 if big else .45),w-.3,h-(1.1 if big else .58),body,10,MUTED)
def arrow(s,x,y,w,color=CYAN):
    line(s,x,y,x+w-.12,y,color,2); sh=s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,Inches(x+w-.18),Inches(y-.08),Inches(.18),Inches(.16)); sh.rotation=90; sh.fill.solid(); sh.fill.fore_color.rgb=color; sh.line.color.rgb=color
def node(s,x,y,w,h,label,color=CYAN,sub=''):
    rect(s,x,y,w,h,PANEL2,color,True); text(s,x+.08,y+.1,w-.16,.3,label,12,WHITE,True,align=PP_ALIGN.CENTER); 
    if sub: text(s,x+.08,y+.43,w-.16,h-.46,sub,8,MUTED,align=PP_ALIGN.CENTER)
def drone(s,x,y,scale=1,color=CYAN):
    # compact technical quadrotor icon
    line(s,x+.15*scale,y+.15*scale,x+1.1*scale,y+.15*scale,color,2); line(s,x+.62*scale,y-.15*scale,x+.62*scale,y+.42*scale,color,2)
    for dx,dy in [(0,0),(1.0,0),(.15,.55),(.85,.55)]:
        c=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x+dx*scale),Inches(y+dy*scale),Inches(.22*scale),Inches(.12*scale)); c.fill.solid(); c.fill.fore_color.rgb=BG; c.line.color.rgb=color
    b=rect(s,x+.46*scale,y+.14*scale,.35*scale,.25*scale,color,color,True); return b
def bullet_list(s,x,y,w,items,size=12,color=WHITE,accent=CYAN,spacing=.34):
    for i,it in enumerate(items):
        yy=y+i*spacing; rect(s,x,yy,.08,.08,accent,accent,True); text(s,x+.18,yy-.05,w-.2,.25,it,size,color)

# 1
s=base('RescueSwarm','Autonomous emergency response platform',1); text(s,.62,1.55,6.8,.75,'Autonomous Multi-Drone\nSearch, Survivor Geolocation & Kit Delivery',27,WHITE,True); text(s,.66,3.2,5.6,.3,'SEARCH  •  DETECT  •  LOCATE  •  DELIVER  •  RETURN',11,CYAN,True,font=MONO)
text(s,.66,4.0,4.7,.5,'One operator. Multiple autonomous vehicles.\nOne coordinated rescue mission.',17,MUTED); pill(s,.66,5.2,1.45,'AUTONOMOUS',CYAN); pill(s,2.25,5.2,1.35,'COLLABORATIVE',BLUE); pill(s,3.75,5.2,1.0,'OFFLINE',GREEN); pill(s,4.95,5.2,1.2,'EDGE AI',ORANGE)
rect(s,7.2,1.25,5.45,5.1,PANEL,GRID,True); text(s,7.55,1.52,4.7,.25,'CONCEPTUAL MISSION VIEW',9,MUTED,True,font=MONO)
for i in range(8): line(s,7.5,2.2+i*.45,12.35,2.2+i*.45,GRID,1)
for i in range(9): line(s,7.5+i*.55,2.2,7.5+i*.55,5.35,GRID,1)
for x,y in [(8.1,2.65),(10.3,3.3),(11.4,2.35)]: drone(s,x,y,.8,CYAN)
for x,y in [(8.5,4.7),(10.0,4.2),(11.5,4.75)]:
    c=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(.18),Inches(.18)); c.fill.solid(); c.fill.fore_color.rgb=RED; c.line.color.rgb=RED
    line(s,x+.09,y+.09,x+.35,y-.18,RED,1)
text(s,7.55,5.65,4.7,.3,'SIMULATION-STYLE CONCEPT • NOT PHYSICAL DEPLOYMENT',8,ORANGE,True,font=MONO,align=PP_ALIGN.CENTER)

# 2
s=base('The Rescue Challenge','Problem',2); text(s,.7,1.55,5.4,.55,'Detection alone is not enough.',24,WHITE,True); text(s,.7,2.2,5.1,.55,'A rescue system must detect, locate, coordinate and respond under time pressure.',14,MUTED)
card(s,.7,3.2,3.7,2.75,'Conventional search','Large areas • manual search\nApproximate visual sightings\nAid arrives after the search',RED)
card(s,4.82,3.2,3.7,2.75,'The bottleneck','One vehicle creates a coverage\nand response constraint.\nConnectivity may be unavailable.',ORANGE)
card(s,8.94,3.2,3.7,2.75,'The requirement','Parallel search\nReliable survivor coordinates\nAid delivery before recovery',CYAN)
arrow(s,4.45,4.55,.3,ORANGE); arrow(s,8.57,4.55,.3,CYAN); text(s,.7,6.4,11.8,.35,'THE OPPORTUNITY  /  TURN A FLEET INTO A COORDINATED RESCUE SYSTEM',11,CYAN,True,font=MONO,align=PP_ALIGN.CENTER)

# 3
s=base('RescueSwarm','Solution',3); text(s,.7,1.42,11.6,.4,'From mission boundary to safely completed response.',18,WHITE,True)
steps=['BOUNDARY','TAKEOFF','PARALLEL\nSEARCH','DETECT +\nCONFIRM','GEOLOCATE','ASSIGN','APPROACH','DELIVER','RESUME','RTL / LAND']
for i,st in enumerate(steps):
    x=.7+i*1.23; node(s,x,2.45,1.02,1.08,st,CYAN if i<5 else ORANGE if i<8 else GREEN); 
    if i<9: arrow(s,x+1.04,2.98,.18,MUTED)
text(s,.8,4.4,11.5,.35,'One operator supervises. The software executes the mission.',22,WHITE,True,align=PP_ALIGN.CENTER)
card(s,1.55,5.25,3.15,1.0,'Autonomy','After START, mission logic drives the fleet.',BLUE)
card(s,5.1,5.25,3.15,1.0,'Coordination','Shared targets, ownership and reallocation.',PURPLE)
card(s,8.65,5.25,3.15,1.0,'Safety','Native RTL/LAND, geofence and recall paths.',GREEN)

# 4
s=base('Designed Around the Full Rescue Mission','Competition scenario',4); text(s,.7,1.4,6.3,.35,'A software-scale rehearsal of the operating envelope.',18,WHITE,True); bullet_list(s,.8,2.05,5.4,['Search area up to approximately 10 hectares','Up to 10 survivors','2+ collaborative autonomous drones','Local / offline communication','One operator / one GCS','Autonomous operation after START'],13,WHITE,CYAN,.48)
rect(s,7.0,1.5,5.55,4.7,PANEL,GRID,True); text(s,7.3,1.78,4.8,.25,'PARTITIONED SEARCH REGION',9,MUTED,True,font=MONO)
for i,col in enumerate([BLUE,CYAN,PURPLE]):
    rect(s,7.35+i*1.62,2.35,1.52,2.8,RGBColor(23+i*5,43+i*3,66+i*4),col,True); text(s,7.45+i*1.62,2.55,1.3,.25,f'DRONE 0{i+1}',10,col,True,font=MONO,align=PP_ALIGN.CENTER); drone(s,7.8+i*1.62,3.1,.6,col)
for x,y in [(7.75,5.15),(9.3,4.8),(10.8,5.2),(11.6,3.8)]:
 c=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(.16),Inches(.16)); c.fill.solid(); c.fill.fore_color.rgb=RED; c.line.color.rgb=RED
text(s,.7,6.5,11.8,.3,'DESIGNED TARGET • NOT CLAIMED AS COMPLETED PHYSICAL TESTING',10,ORANGE,True,font=MONO,align=PP_ALIGN.CENTER)

# 5
s=base('End-to-End System Architecture','Architecture',5)
text(s,.72,1.35,3.6,.25,'GROUND / OPERATOR',10,CYAN,True,font=MONO); text(s,.72,4.2,3.6,.25,'AIRBORNE',10,CYAN,True,font=MONO); text(s,9.55,1.35,3,.25,'AUTOPILOT',10,CYAN,True,font=MONO)
nodes=[('Raspberry Pi 4','GCS'),('React + TypeScript + Vite','MapLibre • Three.js'),('FastAPI','WebSocket'),('PostgreSQL + PostGIS','Fleet coordinator')]
for i,(a,b) in enumerate(nodes): node(s,.72+i*2.15,1.75,1.82,.92,a,BLUE,b)
for i in range(3): arrow(s,2.58+i*2.15,2.21,.42,BLUE)
for i,(a,b) in enumerate([('Raspberry Pi 5','Camera adapter'),('Vision pipeline','DetectorBackend'),('Target manager','Mission manager'),('Fleet gateway','MAVLink')]): node(s,.72+(i%4)*2.15,4.55,1.82,.92,a,PURPLE,b)
for i in range(3): arrow(s,2.58+i*2.15,5.01,.42,PURPLE)
node(s,9.55,1.75,2.9,1.05,'Pixhawk / ArduCopter',ORANGE,'GNSS • IMU • EKF • RTL'); node(s,9.55,3.35,2.9,1.05,'AUTO / GUIDED / RTL',GREEN,'Geofence • failsafes'); node(s,9.55,4.95,2.9,1.05,'Payload servo + motors',RED,'MAVLink command path')
line(s,8.95,2.2,9.55,2.2,CYAN,2); line(s,8.95,5.0,9.55,5.0,CYAN,2); text(s,4.35,6.45,4.5,.3,'LOCAL NETWORK ONLY  •  MAVLink  •  NO CLOUD DEPENDENCY',10,MUTED,True,font=MONO,align=PP_ALIGN.CENTER)

# 6
s=base('Autonomous Rescue State Machine','Autonomy',6); states=['READY','TAKEOFF','SEARCH','CANDIDATE','CONFIRM','GEOLOCATE','ASSIGN','APPROACH','REACQUIRE','ALIGN','DELIVER','VERIFY','RESUME']
for i,st in enumerate(states):
    x=.65+(i%7)*1.78; y=1.75+(i//7)*1.4; node(s,x,y,1.42,.65,st,CYAN if i<6 else ORANGE if i<11 else GREEN)
    if i%7<6: arrow(s,x+1.45,y+.33,.22,MUTED)
    if i==6: line(s,1.36,2.4,1.36,3.15,MUTED,1)
for i,st in enumerate(['MISSION COMPLETE','RTL','LAND','DISARM']): node(s,1.25+i*2.75,5.05,2.15,.78,st,GREEN)
arrow(s,11.3,3.88,.45,GREEN); text(s,.8,6.35,11.8,.32,'SAFETY PATHS: BATTERY • LINK • EKF • GEOFENCE • EMERGENCY RECALL  →  RTL / LAND',10,RED,True,font=MONO,align=PP_ALIGN.CENTER)

# 7
s=base('Fleet Intelligence','Multi-drone collaboration',7); text(s,.7,1.4,6,.35,'Central coordination. Independent vehicle agents.',19,WHITE,True); card(s,.7,2.05,3.3,3.3,'FleetMissionCoordinator','Deterministic partitioning\nShared target registry\nDuplicate merge\nTransactional assignment\nEmergency recall',CYAN); card(s,4.35,2.05,3.3,3.3,'VehicleRescueAgent','Independent autonomy\nOwn mission state\nOwn health and link\nDelivery ownership\nSafe recovery paths',PURPLE)
rect(s,8.05,1.65,4.5,4.0,PANEL,GRID,True); text(s,8.35,1.92,3.9,.25,'SHARED SEARCH MAP',9,MUTED,True,font=MONO)
for i,col in enumerate([BLUE,CYAN,PURPLE]): rect(s,8.4+i*1.25,2.45,1.15,2.35,RGBColor(18,32,54),col,True); drone(s,8.65+i*1.25,2.8,.48,col); text(s,8.48+i*1.25,4.45,1,.2,f'SECTOR {i+1}',8,col,True,font=MONO,align=PP_ALIGN.CENTER)
for x,y in [(8.75,4.0),(10.1,3.5),(11.4,4.15)]: c=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(.14),Inches(.14)); c.fill.solid(); c.fill.fore_color.rgb=RED; c.line.color.rgb=RED
text(s,.75,6.35,11.7,.3,'VEHICLE LOSS  →  REMAINING-AREA REALLOCATION  →  NO DUPLICATE DELIVERY',10,CYAN,True,font=MONO,align=PP_ALIGN.CENTER)

# 8
s=base('Pixel-to-Survivor Pipeline','Vision',8); stages=[('CAMERA','FramePacket'),('BOUNDED\nQUEUE','Drop-old'),('YOLO\nBACKEND','Person baseline'),('CANDIDATE','Confidence'),('TEMPORAL\nCONFIRM','Multi-view'),('CANONICAL\nTARGET','Registry')]
for i,(a,b) in enumerate(stages): node(s,.65+i*2.1,2.25,1.7,1.25,a,PURPLE,b); 
for i in range(5): arrow(s,2.38+i*2.1,2.88,.3,MUTED)
card(s,1.0,4.45,3.3,1.05,'Engineering guardrail','Bounded queues prevent unlimited frame backlog.',CYAN); card(s,4.95,4.45,3.3,1.05,'Model posture','Generic person detector is current baseline.',ORANGE); card(s,8.9,4.45,3.3,1.05,'Hardware path','Monochrome can replicate to 3 channels for RGB-shaped models.',GREEN)
text(s,.75,6.35,11.7,.3,'ARCHITECTURE SUPPORTS A FUTURE CUSTOM SURVIVOR DETECTOR  •  NO FINAL PHYSICAL ACCURACY CLAIM',9,MUTED,True,font=MONO,align=PP_ALIGN.CENTER)

# 9
s=base('Detection ≠ Location','Pose-aware geolocation',9); text(s,.7,1.4,5.4,.38,'A bounding box is not a latitude / longitude.',20,WHITE,True); text(s,.7,1.92,5.35,.55,'The system projects a camera ray through the detected pixel and intersects it with the ground model.',13,MUTED)
rect(s,7.0,1.25,5.45,4.95,PANEL,GRID,True); line(s,7.4,5.45,12.0,5.45,ORANGE,2); text(s,7.5,5.6,4.2,.25,'GROUND PLANE',9,ORANGE,True,font=MONO)
drone(s,9.25,1.75,1.0,CYAN); line(s,9.76,2.4,8.65,5.45,CYAN,2); line(s,9.76,2.4,10.65,5.45,CYAN,2); line(s,9.76,2.4,9.92,5.45,CYAN,2); text(s,9.95,1.52,1.2,.25,'ALTITUDE',9,WHITE,True,font=MONO)
c=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(9.8),Inches(4.93),Inches(.18),Inches(.18)); c.fill.solid(); c.fill.fore_color.rgb=RED; c.line.color.rgb=RED; text(s,10.05,4.85,1.4,.3,'SURVIVOR',9,RED,True)
bullet_list(s,.8,3.0,5.2,['Camera intrinsics + lens model','Pixel + mounting position / rotation','Drone position + altitude','Roll, pitch, yaw / quaternion','Frame timestamp + ground model'],13,WHITE,CYAN,.48)

# 10
s=base('Pixel → World Coordinates','Spatial transforms',10); text(s,.7,1.4,11.5,.35,'The transform chain is explicit, testable and quaternion-authoritative.',18,WHITE,True)
for i,(lab,sub,col) in enumerate([('PIXEL','undistorted ray',PURPLE),('CAMERA OPTICAL','x right • y down • z forward',PURPLE),('BODY FRD','x forward • y right • z down',BLUE),('NED','world navigation frame',CYAN),('MISSION ENU','local map frame',CYAN),('WGS84','latitude / longitude',GREEN)]):
    x=.65+i*2.1; node(s,x,2.35,1.65,1.05,lab,col,sub)
    if i<5: arrow(s,x+1.7,2.87,.3,MUTED)
text(s,.9,4.35,11.2,.6,'T_world_camera  =  T_world_body  ×  T_body_camera',24,WHITE,True,font=MONO,align=PP_ALIGN.CENTER)
text(s,.9,5.25,11.2,.35,'pixel → camera ray → body ray → world ray → ground intersection → ENU → latitude / longitude',12,MUTED,align=PP_ALIGN.CENTER)
pill(s,4.25,6.0,1.45,'QUATERNION',CYAN); text(s,5.85,6.03,3.5,.23,'internal 3D orientation authority',11,MUTED)

# 11
s=base('Real Flight Disturbance Changes Geolocation','Wind / physics',11); text(s,.7,1.4,9.2,.35,'Wind is not a decorative animation.',22,WHITE,True); text(s,.7,1.95,9.5,.32,'Gazebo applies force. ArduPilot responds. The camera pose inherits the aircraft attitude.',13,MUTED)
chain=['WIND / GUST','FORCE','ATTITUDE','CAMERA\nFOOTPRINT','PIXEL','RAY','COORDINATE']
for i,st in enumerate(chain): node(s,.7+i*1.77,3.0,1.4,1.1,st,ORANGE if i<3 else RED if i<5 else CYAN)
for i in range(6): arrow(s,2.15+i*1.77,3.55,.28,MUTED)
card(s,1.1,5.05,4.9,1.0,'CALM','Stable footprint; smaller pose variation.',BLUE); card(s,7.0,5.05,4.9,1.0,'CROSSWIND / GUST','Natural attitude response shifts the projected ray.',ORANGE)
text(s,.7,6.48,11.8,.25,'PHYSICS STRESS TEST  /  EXTREME ~60° PEAK-TO-PEAK ATTITUDE VARIATION IS NOT REPRESENTATIVE FLIGHT VALIDATION',8,ORANGE,True,font=MONO,align=PP_ALIGN.CENTER)

# 12
s=base('The Right Pose at the Right Frame','Time synchronization',12); text(s,.7,1.4,9.8,.35,'Using the latest pose can be wrong during rapid motion.',20,WHITE,True)
for i,(lab,sub,col) in enumerate([('CAMERA FRAME','capture at T',PURPLE),('POSE HISTORY','MAVLink samples',BLUE),('CLOCK MAP','simulation ↔ autopilot',CYAN),('SLERP','quaternion interpolation',GREEN),('POSE @ T','accepted for geometry',ORANGE)]):
    node(s,.75+i*2.45,2.35,2.0,1.08,lab,col,sub)
    if i<4: arrow(s,2.83+i*2.45,2.9,.38,MUTED)
text(s,1.0,4.3,11.2,.55,'unmapped capture timestamp  →  REJECT GEOLOCATION',23,WHITE,True,font=MONO,align=PP_ALIGN.CENTER)
text(s,1.0,5.05,11.2,.5,'The current Phase-2 path fails closed instead of substituting frame receipt time.',14,ORANGE,True,align=PP_ALIGN.CENTER)
pill(s,4.7,6.05,1.25,'HONEST',GREEN); text(s,6.1,6.08,3.0,.25,'safety-oriented engineering',11,MUTED)

# 13
s=base('Spatial Digital Twin','GCS + simulation observability',13); text(s,.7,1.4,8.5,.35,'Two views of reality: autonomy estimate vs simulation ground truth.',18,WHITE,True)
rect(s,.7,2.0,7.25,4.05,PANEL,GRID,True); text(s,1.0,2.25,3.5,.25,'3D MISSION VIEW',9,MUTED,True,font=MONO); line(s,1.15,5.2,7.45,5.2,GRID,1); line(s,1.4,3.0,1.4,5.2,GRID,1); line(s,1.4,5.2,7.25,3.0,GRID,1); drone(s,3.9,2.7,1.0,CYAN); line(s,4.5,3.35,3.5,5.2,CYAN,2); line(s,4.5,3.35,5.4,5.2,CYAN,2); c=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(5.32),Inches(4.8),Inches(.2),Inches(.2)); c.fill.solid(); c.fill.fore_color.rgb=RED; c.line.color.rgb=RED; text(s,5.6,4.76,1.2,.2,'target ray',8,RED,True); text(s,3.85,5.45,2,.2,'GROUND FOOTPRINT',8,ORANGE,True,font=MONO,align=PP_ALIGN.CENTER)
card(s,8.35,2.0,4.2,1.15,'Visualizes','Position • quaternion • paths • frustum • footprint',CYAN); card(s,8.35,3.35,4.2,1.15,'Explains','Target rays • uncertainty • sectors • deliveries',PURPLE); card(s,8.35,4.7,4.2,1.15,'Replays','Mission timeline • engineering inspector • observer',GREEN)
text(s,.7,6.45,11.7,.3,'0.000e+00 m  /  deterministic Python ↔ TypeScript spatial fixture discrepancy',12,GREEN,True,font=MONO,align=PP_ALIGN.CENTER)

# 14
s=base('High-Fidelity Closed-Loop Simulation Lab','Simulation',14); text(s,.7,1.38,10.8,.35,'A configurable digital test range for the full autonomy loop.',18,WHITE,True)
labels=[('SIM CONFIG','drone • camera • environment'),('GAZEBO','physics • wind • rendering'),('ARDUPILOT SITL','EKF • control • MAVLink'),('RESCUESWARM','vision • mission • fleet'),('EVALUATION','truth after prediction')]
for i,(a,b) in enumerate(labels): node(s,.7+i*2.45,2.35,2.0,1.15,a,[BLUE,CYAN,ORANGE,PURPLE,GREEN][i],b); 
for i in range(4): arrow(s,2.8+i*2.45,2.92,.42,MUTED)
card(s,.95,4.45,3.55,1.1,'DRONE CONFIG','Frame • mass • payload • speed • inertia / drag',BLUE); card(s,4.9,4.45,3.55,1.1,'CAMERA CONFIG','Resolution • FPS • FOV • mount • latency / noise',PURPLE); card(s,8.85,4.45,3.55,1.1,'ENVIRONMENT','Wind • gusts • lighting • origin • hidden targets',ORANGE)
text(s,.7,6.45,11.7,.28,'Gazebo Harmonic  /  gz-sim 8.15.0  /  official ardupilot_gazebo integration',10,MUTED,True,font=MONO,align=PP_ALIGN.CENTER)

# 15
s=base('No Ground-Truth Leakage','Simulation integrity',15); text(s,.7,1.42,10.6,.38,'The simulator knows the answer. Autonomy does not.',22,WHITE,True)
rect(s,.75,2.15,4.8,3.15,PANEL,GRID,True); text(s,1.05,2.45,3.8,.25,'SIMULATOR KNOWS',10,ORANGE,True,font=MONO); bullet_list(s,1.05,2.95,3.9,['Actual survivor location','True aircraft pose','True wind','Rendered visibility'],14,WHITE,ORANGE,.48)
rect(s,7.75,2.15,4.8,3.15,PANEL,GRID,True); text(s,8.05,2.45,3.8,.25,'AUTONOMY KNOWS ONLY',10,CYAN,True,font=MONO); bullet_list(s,8.05,2.95,3.9,['Camera pixels','MAVLink / EKF pose','Detector output','Mission state'],14,WHITE,CYAN,.48)
rect(s,5.78,2.78,1.72,1.9,RED,RED,True); text(s,5.92,3.22,1.45,.75,'TRUTH\nFIREWALL',16,WHITE,True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
text(s,.8,5.85,11.7,.55,'Ground truth is used only after prediction for scoring / evaluation.',18,GREEN,True,align=PP_ALIGN.CENTER)

# 16
s=base('What the Virtual Drone Actually Sees','Pixel-only perception',16); text(s,.7,1.38,10.8,.35,'Rendered pixels first. Evaluation truth later.',18,WHITE,True)
for i,(lab,col) in enumerate([('NO PERSON',MUTED),('EDGE ENTRY',BLUE),('VISIBLE',CYAN),('YOLO BOX',PURPLE),('CONFIRMED',GREEN)]):
    x=.7+i*2.45; rect(s,x,2.1,2.0,1.55,PANEL2,GRID,True); text(s,x+.12,2.24,1.75,.2,lab,8,col,True,font=MONO,align=PP_ALIGN.CENTER); line(s,x+.3,3.25,x+1.7,3.25,GRID,1)
    if i>=1: c=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x+.88),Inches(2.6),Inches(.28),Inches(.5)); c.fill.solid(); c.fill.fore_color.rgb=RED; c.line.color.rgb=RED
    if i==3: rect(s,x+.7,2.45,.65,.9,RED,RED)
    if i<4: arrow(s,x+2.06,2.88,.32,MUTED)
text(s,.9,4.25,11.2,.3,'Hidden 3D survivor → rendered camera → YOLO11n ONNX → bbox → confirmation → geolocation',12,MUTED,align=PP_ALIGN.CENTER)
card(s,.75,4.85,3.65,1.15,'SIMULATION DEV EVIDENCE','TP 63 • FP 3 • FN 438\nConditional precision 0.9545',ORANGE); card(s,4.85,4.85,3.65,1.15,'BOUNDING BOX','P50 IoU 0.5762\nP95 IoU 0.6879',PURPLE); card(s,8.95,4.85,3.65,1.15,'QUALIFICATION STATUS','Phase 2 remains under qualification\nframe-time chain incomplete',RED)
text(s,.7,6.45,11.8,.25,'NOT FINAL REAL-WORLD DETECTOR ACCURACY  •  model SHA256: 417c59a7…6f2a5e2d  •  threshold 0.25',8,ORANGE,True,font=MONO,align=PP_ALIGN.CENTER)

# 17
s=base('Software Qualification','Verified status',17); text(s,.7,1.38,10.7,.35,'Software verified. Counts evolve as development progresses.',18,WHITE,True)
card(s,.75,2.0,2.7,1.55,'PRIMARY STAGE 12','328',CYAN,big='328'); card(s,3.7,2.0,2.7,1.55,'LATEST PHASE 2','371 combined',CYAN,big='371'); card(s,6.65,2.0,2.7,1.55,'INTENTIONAL SKIPS','9',ORANGE,big='9'); card(s,9.6,2.0,2.7,1.55,'SPATIAL FIXTURE','0.000e+00 m',GREEN,big='0.000e+00')
text(s,.93,3.72,11,.28,'Python 335 passed / 9 skipped  •  Frontend 36 passed  •  Combined 371 passed / 9 skipped',10,MUTED,font=MONO,align=PP_ALIGN.CENTER)
card(s,.75,4.25,3.65,1.15,'OFFLINE','PASS • zero external runtime dependencies in qualified GCS',GREEN); card(s,4.85,4.25,3.65,1.15,'BUILD','Production frontend build PASS',GREEN); card(s,8.95,4.25,3.65,1.15,'DISPLAY QA','1366×768 • 1536×864 • 1920×1080 • 125% zoom',GREEN)
text(s,.75,6.35,11.8,.28,'DETERMINISTIC REPLAY PASS  •  COMPETITION MODE RESTRICTIONS PASS  •  ZERO HORIZONTAL OVERFLOW',9,GREEN,True,font=MONO,align=PP_ALIGN.CENTER)

# 18
s=base('Multi-Drone Qualification in ArduPilot SITL','SITL / software qualification',18); text(s,.7,1.38,11,.35,'Competition-scale software rehearsal, not physical flight.',18,WHITE,True)
card(s,.75,2.0,2.25,1.35,'VEHICLES','2',BLUE,big='2'); card(s,3.2,2.0,2.25,1.35,'DELIVERIES','4',GREEN,big='4'); card(s,5.65,2.0,2.25,1.35,'COVERAGE','99.7453%',CYAN,big='99.7453%'); card(s,8.1,2.0,2.25,1.35,'INTERVENTIONS','0',GREEN,big='0'); card(s,10.55,2.0,2.0,1.35,'OUTCOME','SUCCESS',GREEN,big='SUCCESS')
text(s,.8,3.58,11.7,.28,'4 synthetic survivor targets • fleet distance 669.67 m • both vehicles RTL / landed / disarmed',10,MUTED,font=MONO,align=PP_ALIGN.CENTER)
card(s,.8,4.15,3.55,1.2,'SCALE','99,584.34 m² ≈ 9.9584 ha boundary',CYAN); card(s,4.9,4.15,3.55,1.2,'REGISTRY','10-survivor registry PASS\n3-vehicle smoke test PASS',PURPLE); card(s,9.0,4.15,3.55,1.2,'RECOVERY','Duplicate prevention PASS\nVehicle-loss reallocation PASS',GREEN)
text(s,.8,6.35,11.7,.28,'ARDUPILOT SITL  /  ArduCopter 4.8.0 DEV  /  commit af2a1baf…78d11a',9,ORANGE,True,font=MONO,align=PP_ALIGN.CENTER)

# 19
s=base('One Operator — Complete Mission Visibility','GCS experience',19); text(s,.7,1.35,10.7,.35,'A single local control surface for setup, supervision and review.',18,WHITE,True)
rect(s,.75,2.0,11.8,4.15,PANEL,GRID,True); rect(s,1.0,2.27,11.3,.45,PANEL2,GRID,True); text(s,1.2,2.38,3,.2,'MISSION STATUS  /  FLEET HEALTH',9,CYAN,True,font=MONO); text(s,10.0,2.38,2,.2,'T+ 08:42',12,WHITE,True,font=MONO,align=PP_ALIGN.RIGHT)
rect(s,1.0,2.95,7.2,2.45,RGBColor(14,28,47),GRID,True); text(s,1.2,3.15,2,.2,'2D / 3D MISSION MAP',8,MUTED,True,font=MONO)
for i in range(6): line(s,1.2,3.55+i*.3,7.85,3.55+i*.3,GRID,1)
for i in range(8): line(s,1.4+i*.8,3.4,1.4+i*.8,5.2,GRID,1)
for x,y,col in [(2.0,4.2,BLUE),(4.2,3.7,CYAN),(6.4,4.5,PURPLE)]: drone(s,x,y,.55,col)
rect(s,8.5,2.95,3.55,2.45,PANEL2,GRID,True); text(s,8.75,3.15,2.8,.2,'FLEET / TARGET CARDS',8,MUTED,True,font=MONO); bullet_list(s,8.75,3.55,2.8,['DRONE 01  SEARCH','DRONE 02  DELIVER','TARGET 04  CONFIRMED','PAYLOAD  3 / 4'],11,WHITE,CYAN,.4)
text(s,1.15,5.7,10.7,.2,'SETUP   MISSION   SIMULATION   ENGINEERING   REVIEW',9,CYAN,True,font=MONO,align=PP_ALIGN.CENTER)

# 20
s=base('Built for Safety and Infrastructure Independence','Safety + offline',20); card(s,.75,1.6,5.55,3.9,'SAFETY BY DESIGN','Native ArduPilot RTL / LAND respected\nGeofence integration\nEmergency recall + safety abort\nHealth freshness + stale telemetry\nBounded queues + durable mission events\nRecovery after service restart\nCritical-failure target / drop gating',RED); card(s,7.05,1.6,5.55,3.9,'OFFLINE BY DESIGN','Local network only\nNo cloud inference\nNo GSM / LTE / 5G dependency\nNo public Wi-Fi\nLocally bundled frontend assets\nLocal database\nOffline runtime compliance verified',CYAN)
text(s,.8,6.25,11.7,.35,'COMPETITION MODE  /  allowed controls: mission load, START, safety abort, emergency recall',12,ORANGE,True,font=MONO,align=PP_ALIGN.CENTER)

# 21
s=base('Hardware Integration','Physical status',21); text(s,.7,1.38,9.5,.35,'The hardware path is defined. Physical validation is still pending.',18,WHITE,True)
parts=[('AIRBORNE','Raspberry Pi 5 • OV9281 mono global shutter',PURPLE),('FLIGHT','Pixhawk 2.4.8 • ArduCopter • M8N-class GNSS',ORANGE),('GROUND','Raspberry Pi 4 • local GCS',BLUE),('PAYLOAD','Emergency kit mechanism • custom airframe',RED)]
for i,(a,b,col) in enumerate(parts): node(s,.75+(i%2)*3.25,2.1+(i//2)*1.35,2.75,.95,a,col,b)
rect(s,7.55,1.8,4.95,3.6,PANEL,GRID,True); text(s,7.85,2.1,4.2,.25,'CURRENT STATUS',10,ORANGE,True,font=MONO); pill(s,7.85,2.62,1.45,'SOFTWARE',GREEN); text(s,9.5,2.65,2.5,.2,'release candidate',11,WHITE); pill(s,7.85,3.25,1.45,'SIMULATION',CYAN); text(s,9.5,3.28,2.5,.2,'high-fidelity active',11,WHITE); pill(s,7.85,3.88,1.45,'PHYSICAL',ORANGE); text(s,9.5,3.91,2.5,.2,'pending / diagnosis',11,WHITE)
bullet_list(s,7.85,4.45,4.0,['Battery packs / safe power','OV9281 hardware troubleshooting','Pixhawk / M8N validation','Camera calibration + payload test','Real flight validation'],10,WHITE,ORANGE,.3)
text(s,.75,6.45,11.7,.25,'PENDING ≠ FAILED  •  no physical-flight or payload-delivery claim is made',10,ORANGE,True,font=MONO,align=PP_ALIGN.CENTER)

# 22
s=base('From Digital Twin to Physical Rescue Flight','Validation roadmap',22); text(s,.7,1.38,11,.35,'Progression is gated by evidence, not by optimism.',18,WHITE,True)
items=[('COMPLETED','Core architecture • planning • single-drone autonomy • fleet SITL • GCS • 3D twin',GREEN),('ACTIVE','Pixel-only perception integration • high-fidelity Gazebo physics',CYAN),('NEXT','Frame-time mapping • sensor-noise realism • full simulated rescue mission',ORANGE),('HARDWARE GATE','Camera calibration • physical geolocation • payload • real multi-drone flight',RED)]
for i,(a,b,col) in enumerate(items): y=2.0+i*.9; rect(s,.8,y,.18,.55,col,col,True); text(s,1.2,y+.03,1.55,.22,a,10,col,True,font=MONO); text(s,2.85,y+.03,9.3,.38,b,12,WHITE)
line(s,.88,2.55,.88,5.55,GRID,2)
text(s,.8,6.35,11.7,.38,'RescueSwarm is designed not merely to detect survivors — but to autonomously search, locate, coordinate, deliver aid and safely complete the mission.',13,WHITE,True,align=PP_ALIGN.CENTER)

# 23
s=base('RescueSwarm','Closing statement',23); text(s,.75,1.75,11.8,.8,'Autonomous Intelligence\nfor Faster Emergency Response',31,WHITE,True,align=PP_ALIGN.CENTER); text(s,.9,3.3,11.4,.35,'Search. Detect. Locate. Deliver. Return.',17,CYAN,True,font=MONO,align=PP_ALIGN.CENTER)
for i,(lab,col) in enumerate([('AUTONOMOUS',CYAN),('COLLABORATIVE',BLUE),('OFFLINE',GREEN),('POSE-AWARE',PURPLE),('VISION-DRIVEN',ORANGE),('SAFETY-ORIENTED',RED)]): pill(s,1.0+i*2.0,4.45,1.75,lab,col)
text(s,.9,6.15,11.4,.3,'SOFTWARE VERIFIED  •  SITL / GAZEBO SIMULATION VERIFIED  •  PHYSICAL HARDWARE PENDING',10,ORANGE,True,font=MONO,align=PP_ALIGN.CENTER)

# speaker note-like footer disclaimer on every slide is intentionally absent to keep slides clean; statuses are on relevant slides.
prs.core_properties.title='RescueSwarm — Autonomous Multi-Drone Search and Emergency Kit Delivery'
prs.core_properties.subject='Competition presentation; software/SITL verified, physical hardware pending'
prs.core_properties.author='RescueSwarm'
path=OUT/'RescueSwarm_Competition_Deck.pptx'; prs.save(path)
print(path)
